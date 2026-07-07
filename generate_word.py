import csv
import sys
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import nsdecls
    from docx.oxml import parse_xml
except ImportError:
    print("python-docx not installed. Please install it with 'pip install python-docx'.")
    sys.exit(1)

doc = Document()

# Add Title
title = doc.add_heading('Checklist Keputusan Owner - Migrasi Data Probetes ERP', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add some intro text
p = doc.add_paragraph('Dokumen ini berisi pertanyaan keputusan untuk ')
p.add_run('Owner').bold = True
p.add_run(', ditulis dengan bahasa sederhana agar mudah dipahami.\n\n')
p.add_run('Petunjuk Pengisian:\n').bold = True
p.add_run('1. Silakan baca pertanyaan pada setiap tabel di bawah ini.\n')
p.add_run('2. Cukup ketik/tuliskan jawaban Anda di kolom yang berwarna kuning (')
p.add_run('Kolom Jawaban Owner').bold = True
p.add_run(').\n')
p.add_run('3. Jika Anda ragu, Anda bisa mengikuti saran pada bagian "Rekomendasi Sistem".\n')

current_section = ""
table = None

with open('Checklist_Owner_Probetes.csv', 'r', encoding='utf-8') as f:
    reader = list(csv.reader(f))
    headers = reader[0]
    
    for row in reader[1:]:
        section = row[0]
        if section != current_section:
            doc.add_paragraph() # Add space before new section
            # Create a new heading for the section
            doc.add_heading(section, level=1)
            
            # Create a table for this section
            table = doc.add_table(rows=1, cols=4)
            table.style = 'Table Grid'
            
            # Set table headers
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'No'
            hdr_cells[1].text = 'Pertanyaan & Detail'
            hdr_cells[2].text = 'Pilihan & Rekomendasi'
            hdr_cells[3].text = 'Jawaban Owner (Isi Di Sini)'
            
            # Make headers bold and add background color
            for cell in hdr_cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.size = Pt(11)
                
                # Add background color to header cells (Blue-ish)
                shading_elm = parse_xml(r'<w:shd {} w:fill="D9E1F2"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(shading_elm)
            
            current_section = section
            
        # Add row to the current table
        row_cells = table.add_row().cells
        
        # No
        row_cells[0].text = row[1]
        
        # Pertanyaan & Detail (Combine Topik and Detail)
        p1 = row_cells[1].paragraphs[0]
        p1.add_run(row[2]).bold = True
        if row[3] and row[3] != '-':
            p1.add_run('\n\nContoh/Detail: ').italic = True
            p1.add_run(row[3])
            
        # Pilihan & Rekomendasi
        p2 = row_cells[2].paragraphs[0]
        p2.add_run('Pilihan:\n').bold = True
        p2.add_run(row[4])
        if row[5] and row[5] != '-':
            p2.add_run('\n\nRekomendasi Sistem:\n').bold = True
            p2.add_run(row[5])
            
        # Jawaban Owner (Empty space to write)
        row_cells[3].text = "\n\n\n" # Add some blank lines for spacing if printed
        
        # Give Jawaban Owner cell a light yellow background to indicate it needs to be filled
        shading_elm = parse_xml(r'<w:shd {} w:fill="FFF2CC"/>'.format(nsdecls('w')))
        row_cells[3]._tc.get_or_add_tcPr().append(shading_elm)
        
doc.save('Checklist_Owner_Probetes.docx')
print("Berhasil membuat Checklist_Owner_Probetes.docx")
